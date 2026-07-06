#!/usr/bin/env python3
"""Build the concise investor briefing (.docx) from the measured data + figures.
Clean, neutral style; short bullets; embedded charts. Run: python analysis/build_docx.py
"""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
CSV = os.path.join(HERE, "measured_convergence_ranking.csv")
OUT = os.path.join(HERE, "UK-Rent-Convergence-Briefing.docx")

NAVY = RGBColor(0x1F, 0x3B, 0x57)
STEEL = RGBColor(0x2C, 0x6E, 0x9E)
GREY = RGBColor(0x60, 0x60, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

df = pd.read_csv(CSV)


def clean(s):
    return (s.replace(" UA", "").replace(", City of", "")
             .replace("North East Lincolnshire", "NE Lincs (Grimsby)")
             .replace("North Lincolnshire", "N Lincs (Scunthorpe)"))


# ---------- styling helpers ----------
def shade(cell, hexfill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(el)


def set_cell(cell, text, bold=False, color=None, size=9, align="left", fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if fill: shade(cell, fill)


def h1(doc, text):
    p = doc.add_paragraph(); p.space_after = Pt(6)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    # bottom border
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in [("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", "2C6E9E")]:
        bottom.set(qn(k), v)
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = STEEL
    return p


def bullet(doc, text, bold_lead=None, sub=False):
    p = doc.add_paragraph(style="List Bullet 2" if sub else "List Bullet")
    p.paragraph_format.space_after = Pt(2)
    if bold_lead:
        r = p.add_run(bold_lead + " "); r.bold = True; r.font.size = Pt(10.5)
    r = p.add_run(text); r.font.size = Pt(10.5)
    return p


def para(doc, text, size=10.5, color=None, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.italic = italic
    if color: r.font.color.rgb = color
    return p


def img(doc, name, width=6.4):
    doc.add_picture(os.path.join(FIG, name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


REGION_ABBR = {"Yorkshire and the Humber": "Yorks & Humber", "North East": "North East",
               "North West": "North West", "West Midlands": "West Midlands",
               "East Midlands": "East Midlands", "South East": "South East",
               "South West": "South West", "East of England": "East of England"}
REGION_FIX = {"Barrow-in-Furness": "North West", "Scarborough": "Yorks & Humber",
              "Northampton": "East Midlands", "Harrogate": "Yorks & Humber",
              "Mendip": "South West"}


def table(doc, headers, rows, widths=None, header_fill="1F3B57", zebra="EEF3F7", fs=9):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"; t.autofit = False; t.allow_autofit = False
    for j, htext in enumerate(headers):
        set_cell(t.rows[0].cells[j], htext, bold=True, color=WHITE, size=fs,
                 align="center", fill=header_fill)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell(cells[j], val, size=fs,
                     align="left" if j == 0 else "center",
                     fill=(zebra if i % 2 else None))
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows:
                r.cells[j].width = Inches(w)
    return t


# ======================= build document =======================
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.8)
    s.left_margin = s.right_margin = Inches(0.85)

# ---- COVER ----
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Rent Convergence in England"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A 1–3 Bed Affordable-Housing Investment Screen"); r.font.size = Pt(15); r.font.color.rgb = STEEL
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Where market rent, LHA, Affordable Rent and social rent line up"); r.italic = True; r.font.size = Pt(12); r.font.color.rgb = GREY
for _ in range(6):
    doc.add_paragraph()
for txt, sz in [("Measured from official open data — 97 local authorities scored", 11),
                ("ONS Private Rental Market Statistics · DWP/VOA LHA · Regulator of Social Housing", 9.5),
                ("July 2026", 10)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.font.size = Pt(sz); r.font.color.rgb = GREY if sz < 11 else NAVY
doc.add_page_break()

# ---- EXECUTIVE SUMMARY ----
h1(doc, "Executive summary")
h2(doc, "What this is")
bullet(doc, "A screen for English areas where all four rent levels for 1–3 bed homes sit close together — lower rent risk, stronger housing-benefit alignment.")
bullet(doc, "Four layers compared: market rent, LHA (benefit cap), Affordable Rent (80% of market), and social rent.")
bullet(doc, "97 local authorities scored from official open data; studios and shared housing excluded.")
h2(doc, "Headline finding")
bullet(doc, "Convergence is concentrated in the low-cost North & Midlands; the least-converged areas are all in the high-rent South.", "Geography:")
bullet(doc, "Every one of the top 25 areas is northern or Midlands; the bottom five are Cambridge, Brighton, Bristol, Exeter, Oxford.")
bullet(doc, "The driver is social-rent attachment — near-uniform social rent only sits close to market rent where market rent is low.", "Why:")
h2(doc, "Top 5 investment picks (convergence + demand)")
bullet(doc, "Teesside (Hartlepool / Redcar / Middlesbrough) · Kingston upon Hull · Doncaster · Sunderland · Stoke-on-Trent.")
img(doc, "fig1_top_areas.png", 6.2)
doc.add_page_break()

# ---- THE IDEA ----
h1(doc, "The convergence idea")
bullet(doc, "LHA close to market rent — benefit-backed demand is well supported.", "Aligned when:")
bullet(doc, "Affordable Rent (80% of market) sits near or below LHA — coverable by a benefit tenant.")
bullet(doc, "Social rent not wildly detached from market rent.")
bullet(doc, "1/2/3-bed rents relatively compressed (low variation).")
para(doc, "Score = 0.40 × LHA-alignment + 0.30 × social-attachment + 0.30 × rent-compression (0–1; higher = tighter).", 9.5, GREY, italic=True)
h2(doc, "The mechanism the data exposes")
bullet(doc, "LHA-to-market alignment is high almost everywhere — LHA is the 30th percentile by design (even Cambridge, Guildford score well).")
bullet(doc, "Social rent barely varies nationally (~£380–£490/mo here) but market rent ranges ~£475–£1,400.")
bullet(doc, "So four-layer convergence needs a low-rent market — the northern signature.", "Result:")
img(doc, "fig2_mechanism.png", 6.2)
doc.add_page_break()

# ---- MEASURED RESULTS: TOP 20 ----
h1(doc, "Measured results — Top 20 areas")
headers = ["#", "Area", "Region", "Mkt 2-bed", "LHA 2-bed", "Social 2-bed", "Score"]
rows = []
for _, r in df.head(20).iterrows():
    rraw = r["region"] if pd.notna(r["region"]) else ""
    region = REGION_ABBR.get(rraw, rraw) if rraw else REGION_FIX.get(r["area"], "—")
    rows.append([int(r["rank"]), clean(r["area"]), region,
                 f"£{int(r['mkt_2'])}", f"£{int(r['lha_2'])}", f"£{int(r['soc_2'])}", f"{r['score']:.3f}"])
table(doc, headers, rows, widths=[0.3, 1.75, 1.35, 0.75, 0.75, 0.85, 0.6], fs=8.5)
para(doc, "Rents matched to the LHA reference window (Oct 2022–Sep 2023); measures structural convergence, "
          "not today's live gap. Full 97-area table in the accompanying spreadsheet/CSV.", 8.5, GREY, italic=True)
doc.add_paragraph()
h2(doc, "Regional pattern")
img(doc, "fig4_regional.png", 5.9)
doc.add_page_break()

# ---- TOP 5 PICKS ----
h1(doc, "Top 5 investment picks")
img(doc, "fig3_four_layers.png", 6.0)
doc.add_paragraph()
picks = [
    ("1 · Teesside — Hartlepool / Redcar / Middlesbrough",
     "Elite convergence (Hartlepool #2, Redcar #3, Middlesbrough #5) with the clearest new-demand story: Teesworks / Freeport and energy investment.",
     "Keep to the regeneration footprint; avoid deprived peripheral estates."),
    ("2 · Kingston upon Hull  (#12)",
     "Top-tier convergence, a large deep market, and a real demand engine — Humber energy cluster, port, university.",
     "Flood-zone and older-stock demand pockets; social rents unusually low."),
    ("3 · Doncaster  (#22)",
     "Strong convergence with a genuine logistics / rail employment base and Sheffield city-region pull.",
     "Post-industrial low-demand micro-markets."),
    ("4 · Sunderland  (#19)",
     "Strong convergence plus Riverside regeneration and automotive employment.",
     "Peripheral estates weaker; single-employer exposure."),
    ("5 · Stoke-on-Trent  (#20)",
     "Strong convergence with regeneration and affordability-led in-migration.",
     "Pockets of low-demand terraced stock; buy in sound sub-markets."),
]
for title, why, risk in picks:
    h2(doc, title)
    bullet(doc, why, "Why:")
    bullet(doc, risk, "Watch:")
para(doc, "Highest pure convergence but thinner demand: Burnley / Hyndburn / Pendle / Blackburn (East Lancs), "
          "Grimsby, Scunthorpe, Rotherham. Blackpool scores well but stays a watch (weak demand, poor stock).",
     9.5, GREY, italic=True)
doc.add_page_break()

# ---- OPERATORS ----
h1(doc, "Operators active in the stream")
para(doc, "Registered providers (RP) and council ALMOs already delivering social & affordable rent in each pick. "
          "RSH grades: G=governance, V=viability, C=consumer (1 = strongest). * = widely-reported figure.", 9.5, GREY, italic=True)
op_headers = ["Area", "Organisation", "Type", "Homes", "RSH (G/V/C)"]
op_rows = [
    ["Teesside", "Thirteen Group", "RP", "~34,000", "G1 / V1 / C1"],
    ["Teesside", "Beyond Housing", "RP", "~15,350", "G1 / V1 / —"],
    ["Teesside", "North Star Housing", "RP", "~4,000", "G1 / V1 / —"],
    ["Hull", "Riverside", "RP", "~75,000*", "G1 / V2 / —"],
    ["Hull", "Sanctuary", "RP", "~120,000*", "G1 / V2 / C2"],
    ["Hull", "Pickering & Ferens Homes", "RP", "~1,400", "n/v"],
    ["Doncaster", "St Leger Homes of Doncaster", "ALMO", "~20,000", "— (council-held)"],
    ["Doncaster", "Together Housing", "RP", "~36,000", "G1 / V2 / —"],
    ["Doncaster", "Housing 21", "RP", "~23,300", "G1 / V1 / C1"],
    ["Sunderland", "Gentoo Group", "RP", "~29,000", "G1 / V2 / C1"],
    ["Stoke-on-Trent", "Aspire Housing", "RP", "~9,300", "G1 / V2 / C1"],
    ["Stoke-on-Trent", "Honeycomb Group (Staffs Housing)", "RP", "~3,118", "n/v"],
    ["Stoke-on-Trent", "EPIC Housing", "RP", "~1,400", "n/v"],
]
table(doc, op_headers, op_rows, widths=[1.0, 2.35, 0.65, 1.0, 1.3], fs=9)
para(doc, "Maps who is active — not an endorsement or a claim of current tenders. Verify current pipeline/JV appetite "
          "and re-check RSH judgements before acting.", 8.5, GREY, italic=True)
doc.add_page_break()

# ---- METHODOLOGY & CAVEATS ----
h1(doc, "Method, data & caveats")
h2(doc, "Data sources (official, open)")
bullet(doc, "Market rent — ONS/VOA Private Rental Market Statistics: median monthly rent by bedroom, by local authority (Oct 2022–Sep 2023).")
bullet(doc, "LHA — DWP/VOA April-2024 rates (frozen), 1/2/3-bed, by BRMA.")
bullet(doc, "Social rent — Regulator of Social Housing 2024/25: per-LA, per-bedsize general-needs rent.")
bullet(doc, "Affordable Rent — derived as 80% of market (statutory definition).")
h2(doc, "How it was built")
bullet(doc, "Three layers joined per area; convergence score computed and ranked (see score formula).")
bullet(doc, "92 of 97 areas use local social rent; 5 use the England average as fallback.")
bullet(doc, "LA→BRMA geography: 56 exact matches + 37 web-verified + 4 genuinely multi-BRMA (flagged).")
h2(doc, "Caveats")
bullet(doc, "Structural, not live — rents matched to the frozen-LHA window; today's market-vs-LHA gap is wider.")
bullet(doc, "A convergence screen, not a demand / yield / asset-quality model — pair with local demand work.")
bullet(doc, "Small-sample areas flagged in the data; treat cautiously.")
bullet(doc, "97 areas scored (not all ~300) — full national coverage needs the official LA→BRMA lookup.")
h2(doc, "Sources")
para(doc, "ONS Private Rental Market Statistics · DWP/VOA Local Housing Allowance rates · Regulator of Social Housing, "
          "Registered provider social housing stock and rents 2024/25 · RSH regulatory judgements (provider grades). "
          "All Open Government Licence. Reproducible via run_measured_ranking.py.", 9, GREY)

doc.save(OUT)
print("wrote", OUT)
